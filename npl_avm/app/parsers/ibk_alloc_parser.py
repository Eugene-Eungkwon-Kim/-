"""
IBK Pre Allocation 파서 ("4. 담보(물건)" 시트)
파일: ★ IBK XXXX-XP A Pre Allocation.xlsx
시트: "4. 담보(물건)"
헤더: Row 8(Field N) + Row 9(한글) → Row 9 사용
특이사항:
  - property_serial이 "R-001_1" 형태로 이미 고유 (debtor prefix 불필요)
  - 감정평가 개별 항목은 "감정평가액\n합계" 기준 offset
"""
import re
from pathlib import Path
import openpyxl

from app.parsers.ibk_parser import IBKParser, C1_HEADER_KEYWORDS
from app.parsers.base_parser import DealRecord, DebtorRecord, PropertyRecord

IBK_ALLOC_KEYWORDS = {
    "일련번호": "serial_no",
    "Pool": "pool_class",
    "자산구분": "debtor_type",
    "차주일련번호": "debtor_serial",
    "차주명": "debtor_name",
    "물건번호": "property_serial",
    "담보소재지1": "address_sido",
    "담보소재지2": "address_sigungu",
    "담보소재지3": "address_dong",
    "담보소재지4": "address_detail",
    "Property Type": "property_type",
    "대지면적": "land_area",
    "건물면적": "building_area",
    "환산후채권최고액": "mortgage_amount_krw",
    "법정선순위 합계": "senior_burden_total",
    "경매\n개시여부": "is_filed",
    "경매신청기관\n(모사건)": "creditor",
    "경매개시일자\n(모사건)": "filing_date",
    "경매 관할법원\n(모사건)": "court",
    "경매사건번호\n(모사건)": "case_number",
    "배당요구종기일\n(모사건)": "demand_deadline",
    "경매청구금액\n(모사건)": "claim_amount",
    "최초경매기일": "first_auction_date",
    "최종경매회차": "lapse_count_raw",
    "최종경매기일": "final_auction_date",
    "최종경매결과": "final_result",
    "낙찰금액": "hammer_price",
    "최종경매일\n최저입찰금액": "final_min_bid",
    "차기경매기일": "next_auction_date",
    "차기경매기일의 최저입찰금액": "next_min_bid",
    "감정평가서 종류": "appraisal_type",
    "감정평가일자": "appraisal_date",
    "감정평가기관": "appraiser",
    "감정평가액\n합계": "total_value",
    "KB아파트시세": "kb_market_price",
}


class IBKAllocParser(IBKParser):

    def parse(self, file_path: str) -> DealRecord:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        deal = self._parse_deal_meta(wb, file_path)
        for sn in wb.sheetnames:
            if "담보(물건)" in sn or "담보 (물건)" in sn:
                self._parse_alloc_sheet(wb[sn], deal)
                break
        wb.close()
        return deal

    def _parse_deal_meta(self, wb, file_path: str) -> DealRecord:
        deal = DealRecord(source_file=file_path, financial_institution="IBK")
        # 시트 첫 행에서 딜명 추출
        for sn in wb.sheetnames:
            if "담보(물건)" in sn:
                ws = wb[sn]
                rows = list(ws.iter_rows(min_row=1, max_row=5, values_only=True))
                for row in rows:
                    for cell in row:
                        if cell and "IBK" in str(cell) and "Program" in str(cell):
                            deal.deal_name = str(cell).strip()
                            break
                break
        if not deal.deal_name:
            stem = Path(file_path).stem
            m = re.search(r"IBK[\s_]*(\d{4}-\d+[\w.]*)", stem, re.IGNORECASE)
            deal.deal_name = f"IBK {m.group(1)} Program" if m else stem[:60]

        # Pool 추출
        pm = re.search(r"Pool\s*([A-Z])", deal.deal_name, re.IGNORECASE)
        deal.pool_name = f"Pool {pm.group(1).upper()}" if pm else "Pool A"
        return deal

    def _build_col_map(self, ws) -> dict[str, int]:
        col_map: dict[str, int] = {}
        all_kw = {**IBK_ALLOC_KEYWORDS, **C1_HEADER_KEYWORDS}

        for row in ws.iter_rows(min_row=6, max_row=11, values_only=True):
            for idx, cell in enumerate(row):
                if cell is None:
                    continue
                cell_str = str(cell).strip()
                for kw, field in all_kw.items():
                    if kw in cell_str and field not in col_map:
                        col_map[field] = idx

        # 감정평가 개별 항목: 합계 기준 offset
        if "total_value" in col_map:
            tc = col_map["total_value"]
            col_map.setdefault("outside_value", tc - 1)
            col_map.setdefault("machine_value", tc - 2)
            col_map.setdefault("building_value", tc - 3)
            col_map.setdefault("land_value", tc - 4)

        return col_map

    def _find_data_start(self, ws) -> int:
        for row_idx, r_obj in enumerate(ws.iter_rows(min_row=10, max_row=20), start=10):
            if len(r_obj) > 1 and r_obj[1].value is not None:
                try:
                    int(float(str(r_obj[1].value).replace(",", "")))
                    return row_idx
                except (ValueError, TypeError):
                    pass
        return 10

    def _parse_alloc_sheet(self, ws, deal: DealRecord):
        debtors: dict[str, DebtorRecord] = {}
        col_map = self._build_col_map(ws)
        data_start = self._find_data_start(ws)

        for row in ws.iter_rows(min_row=data_start, values_only=True):
            if not any(row):
                continue
            if len(row) < 2 or row[1] is None:
                continue

            r = {field: (row[idx] if idx < len(row) else None)
                 for field, idx in col_map.items()}
            for k in ("currency", "debtor_serial", "debtor_name", "debtor_type", "pool_class"):
                r.setdefault(k, None)

            debtor_serial = str(r["debtor_serial"] or "").strip()
            if not debtor_serial:
                continue

            if debtor_serial not in debtors:
                debtors[debtor_serial] = DebtorRecord(
                    debtor_serial=debtor_serial,
                    debtor_name=str(r["debtor_name"] or "").strip(),
                    debtor_type=str(r["debtor_type"] or "Regular").strip(),
                    pool_class=str(r["pool_class"] or "A").strip(),
                )

            # property_serial은 이미 "R-001_1" 형태로 고유
            prop_serial = str(r.get("property_serial") or "").strip()
            if not prop_serial:
                prop_serial = debtor_serial

            prop = PropertyRecord(
                property_serial=prop_serial,
                address_sido=str(r.get("address_sido") or "").strip(),
                address_sigungu=str(r.get("address_sigungu") or "").strip(),
                address_dong=str(r.get("address_dong") or "").strip(),
                address_detail=str(r.get("address_detail") or "").strip(),
                property_type=str(r.get("property_type") or "").strip(),
                land_area=self._clean_float(r.get("land_area")),
                building_area=self._clean_float(r.get("building_area")),
                currency="KRW",
                mortgage_amount=self._clean_amount(r.get("mortgage_amount_krw")),
                senior_burden_total=self._clean_amount(r.get("senior_burden_total")),
                kb_market_price=self._clean_amount(r.get("kb_market_price")),
                appraisal_type=str(r.get("appraisal_type") or "").strip(),
                appraisal_date=self._clean_date(r.get("appraisal_date")),
                appraiser=str(r.get("appraiser") or "").strip(),
                land_value=self._clean_amount(r.get("land_value")),
                building_value=self._clean_amount(r.get("building_value")),
                machine_value=self._clean_amount(r.get("machine_value")),
                outside_value=self._clean_amount(r.get("outside_value")),
                total_value=self._clean_amount(r.get("total_value")),
                is_filed=self._parse_filed(r.get("is_filed")),
                court=str(r.get("court") or "").strip(),
                creditor=str(r.get("creditor") or "").strip(),
                case_number=str(r.get("case_number") or "").strip(),
                filing_date=self._clean_date(r.get("filing_date")),
                demand_deadline=self._clean_date(r.get("demand_deadline")),
                claim_amount=self._clean_amount(r.get("claim_amount")),
                first_auction_date=self._clean_date(r.get("first_auction_date")),
                lapse_count=self._parse_lapse(r.get("lapse_count_raw")),
                final_result=str(r.get("final_result") or "").strip(),
                final_auction_date=self._clean_date(r.get("final_auction_date")),
                next_auction_date=self._clean_date(r.get("next_auction_date")),
                hammer_price=self._clean_amount(r.get("hammer_price")),
                final_min_bid=self._clean_amount(r.get("final_min_bid")),
                next_min_bid=self._clean_amount(r.get("next_min_bid")),
            )
            prop.property_category = self._infer_category(prop.property_type)
            prop.address_full = " ".join(filter(None, [
                prop.address_sido, prop.address_sigungu,
                prop.address_dong, prop.address_detail,
            ]))

            debtors[debtor_serial].properties.append(prop)

        deal.debtors = list(debtors.values())
