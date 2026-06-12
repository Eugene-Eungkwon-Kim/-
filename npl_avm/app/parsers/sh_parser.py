"""
SH(서울주택도시공사) Data Disk 파서 (IBKParser 상속)
시트: Sheet C-1(물건정보)
헤더: Row 10 (키워드 자동감지, C1_FIXED 사용 안 함)
특이사항:
  - IBK와 컬럼 구조 동일하지만 col6 비어있고 col7=물건번호
  - 시트명에 "(물건정보)" 포함
"""
import re
from pathlib import Path
import openpyxl

from app.parsers.ibk_parser import IBKParser, C1_HEADER_KEYWORDS
from app.parsers.base_parser import DealRecord

SH_KEYWORDS = {
    "일련번호": "serial_no",
    "자산유형": "debtor_type",
    "차주일련번호": "debtor_serial",
    "차주명": "debtor_name",
    "물건번호": "property_serial",
    "담보소재지1": "address_sido",
    "담보소재지2": "address_sigungu",
    "담보소재지3": "address_dong",
    "담보소재지4": "address_detail",
    "Property Type": "property_type",
    "대지면적": "land_area",
    "건물면적": "building_area",
    "환산후 Property별\n근저당권설정액": "mortgage_amount_krw",
    "환산후 Property별": "mortgage_amount_krw",
    "Property별 \n선순위 설정순위": "senior_rank",
    "선순위\n가압류/압류/가처분 유무": "has_provisional_seizure",
    "선순위\n가압류/압류/가처분 금액": "provisional_seizure_amount",
    "유치권 유무": "has_lien",
    "유치권 신고금액": "lien_amount",
    "선순위\n소액보증금 (주택)": "small_deposit_housing",
    "선순위\n소액보증금 (상가)": "small_deposit_commercial",
    "선순위\n임차보증금 (주택)": "lease_deposit_housing",
    "선순위\n임차보증금 (상가)": "lease_deposit_commercial",
    "선순위 임금채권": "wage_claim",
    "선순위 당해세": "current_tax",
    "선순위 조세채권": "tax_claim",
    "합계": "senior_burden_total",
    "감정평가구분": "appraisal_type",
    "감정평가일자": "appraisal_date",
    "감정평가기관": "appraiser",
    "토지 감정평가액": "land_value",
    "토지감정평가액": "land_value",
    "건물 감정평가액": "building_value",
    "기계 감정평가액": "machine_value",
    "제시외": "outside_value",
    "감정평가액합계": "total_value",
    "KB아파트시세": "kb_market_price",
    "경매개시": "is_filed",
    "경매 관할법원": "court",
    "경매신청기관\n(모사건)": "creditor",
    "경매개시일자\n(모사건)": "filing_date",
    "경매사건번호\n(모사건)": "case_number",
    "배당요구종기일\n(모사건)": "demand_deadline",
    "청구금액\n(모사건)": "claim_amount",
    "최초법사가": "first_legal_price",
    "최초경매기일": "first_auction_date",
    "최종경매회차": "lapse_count_raw",
    "최종경매결과": "final_result",
    "최종경매기일": "final_auction_date",
    "차기경매기일": "next_auction_date",
    "낙찰금액": "hammer_price",
    "최종경매일의 \n최저입찰금액": "final_min_bid",
    "차후예정경매일의 \n최저입찰금액": "next_min_bid",
}


class SHParser(IBKParser):

    def parse(self, file_path: str) -> DealRecord:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        deal = self._parse_deal_meta(wb, file_path)
        # SH의 시트명: "Sheet C-1(물건정보)"
        for sn in wb.sheetnames:
            if "Sheet C-1" in sn:
                self._parse_c1(wb[sn], deal)
                break
        wb.close()
        return deal

    def _parse_deal_meta(self, wb, file_path: str) -> DealRecord:
        deal = DealRecord(source_file=file_path, financial_institution="SH")
        stem = Path(file_path).stem
        m = re.search(r"SH[\s_]*(\d{4}-\d+[\w.]*)", stem, re.IGNORECASE)
        if m:
            deal.deal_name = f"SH {m.group(1)} Program"
        else:
            deal.deal_name = stem[:60]
        deal.pool_name = "Pool A"
        return deal

    def _build_col_map(self, ws) -> dict[str, int]:
        """C1_FIXED 무시, SH_KEYWORDS + C1_HEADER_KEYWORDS로만 매핑"""
        col_map: dict[str, int] = {}
        all_keywords = {**SH_KEYWORDS, **C1_HEADER_KEYWORDS}

        for row in ws.iter_rows(min_row=7, max_row=13, values_only=True):
            for idx, cell in enumerate(row):
                if cell is None:
                    continue
                cell_str = str(cell).strip()
                for keyword, field in all_keywords.items():
                    if keyword in cell_str and field not in col_map:
                        col_map[field] = idx
            if len(col_map) >= 30:
                break

        return col_map

    def _parse_c1(self, ws, deal: DealRecord):
        """C1_FIXED 없이 키워드 맵만 사용"""
        from app.parsers.base_parser import DebtorRecord, PropertyRecord
        debtors: dict[str, DebtorRecord] = {}
        col_map = self._build_col_map(ws)

        data_start = self._find_data_start(ws)

        for row in ws.iter_rows(min_row=data_start, values_only=True):
            if not any(row):
                continue
            if len(row) < 2 or row[1] is None:
                continue

            r = {field: (row[idx] if idx < len(row) else None)
                 for field, idx in col_map.items()}
            for k in ("currency", "debtor_serial", "debtor_name", "debtor_type", "pool_class"):
                r.setdefault(k, None)

            debtor_serial = str(r["debtor_serial"] or "").strip()
            if not debtor_serial:
                continue

            if debtor_serial not in debtors:
                d = DebtorRecord(
                    debtor_serial=debtor_serial,
                    debtor_name=str(r["debtor_name"] or "").strip(),
                    debtor_type=str(r["debtor_type"] or "Regular").strip(),
                    pool_class="A",
                )
                debtors[debtor_serial] = d

            raw_serial = str(r.get("property_serial") or "").strip()
            prop_serial = f"{debtor_serial}-{raw_serial}" if raw_serial else debtor_serial

            prop = PropertyRecord(
                property_serial=prop_serial,
                address_sido=str(r.get("address_sido") or "").strip(),
                address_sigungu=str(r.get("address_sigungu") or "").strip(),
                address_dong=str(r.get("address_dong") or "").strip(),
                address_detail=str(r.get("address_detail") or "").strip(),
                property_type=str(r.get("property_type") or "").strip(),
                land_area=self._clean_float(r.get("land_area")),
                building_area=self._clean_float(r.get("building_area")),
                currency=str(r.get("currency") or "KRW").strip(),
                mortgage_amount=self._clean_amount(r.get("mortgage_amount_krw")),
                senior_mortgage_amount=self._clean_amount(r.get("senior_mortgage_krw")),
                has_provisional_seizure=self._parse_yn(r.get("has_provisional_seizure")),
                provisional_seizure_amount=self._clean_amount(r.get("provisional_seizure_amount")),
                has_lien=self._parse_yn(r.get("has_lien")),
                lien_amount=self._clean_amount(r.get("lien_amount")),
                small_deposit_housing=self._clean_amount(r.get("small_deposit_housing")),
                small_deposit_commercial=self._clean_amount(r.get("small_deposit_commercial")),
                lease_deposit_housing=self._clean_amount(r.get("lease_deposit_housing")),
                lease_deposit_commercial=self._clean_amount(r.get("lease_deposit_commercial")),
                wage_claim=self._clean_amount(r.get("wage_claim")),
                current_tax=self._clean_amount(r.get("current_tax")),
                tax_claim=self._clean_amount(r.get("tax_claim")),
                senior_burden_total=self._clean_amount(r.get("senior_burden_total")),
                kb_market_price=self._clean_amount(r.get("kb_market_price")),
                appraisal_type=str(r.get("appraisal_type") or "").strip(),
                appraisal_date=self._clean_date(r.get("appraisal_date")),
                appraiser=str(r.get("appraiser") or "").strip(),
                land_value=self._clean_amount(r.get("land_value")),
                building_value=self._clean_amount(r.get("building_value")),
                machine_value=self._clean_amount(r.get("machine_value")),
                outside_value=self._clean_amount(r.get("outside_value")),
                total_value=self._clean_amount(r.get("total_value")),
                is_filed=self._parse_filed(r.get("is_filed")),
                court=str(r.get("court") or "").strip(),
                creditor=str(r.get("creditor") or "").strip(),
                case_number=str(r.get("case_number") or "").strip(),
                filing_date=self._clean_date(r.get("filing_date")),
                demand_deadline=self._clean_date(r.get("demand_deadline")),
                claim_amount=self._clean_amount(r.get("claim_amount")),
                first_legal_price=self._clean_amount(r.get("first_legal_price")),
                first_auction_date=self._clean_date(r.get("first_auction_date")),
                lapse_count=self._parse_lapse(r.get("lapse_count_raw")),
                final_result=str(r.get("final_result") or "").strip(),
                final_auction_date=self._clean_date(r.get("final_auction_date")),
                next_auction_date=self._clean_date(r.get("next_auction_date")),
                hammer_price=self._clean_amount(r.get("hammer_price")),
                final_min_bid=self._clean_amount(r.get("final_min_bid")),
                next_min_bid=self._clean_amount(r.get("next_min_bid")),
            )
            prop.property_category = self._infer_category(prop.property_type)
            prop.address_full = " ".join(filter(None, [
                prop.address_sido, prop.address_sigungu,
                prop.address_dong, prop.address_detail,
            ]))

            debtors[debtor_serial].properties.append(prop)

        deal.debtors = list(debtors.values())
