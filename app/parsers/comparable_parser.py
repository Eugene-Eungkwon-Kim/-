"""
거래사례 정밀 시트 파서
파일: IBK 거래사례/IBK-R-XXX_N ...xlsx, HANA 거래사례/hana-A-R-XXX_N ...xlsx
시트: "정밀"

구조 (R1~R13):
  R1 : 구분       | 본건(C4) | 거래사례1(C6) | 거래사례2(C8) | ...
  R2 : 소재지
  R3 : 용도지역
  R4 : 주용도
  R5 : 거래금액    (본건=감정가, 사례=실거래가)
  R6 : 거래시점/기준시점
  R7 : 일괄단가(원/py)
  R8 : 개별지가(원/㎡)
  R9 : 토지(㎡)
  R10: 건물(㎡)
  R11: 사용승인일
  R12: 주구조
  R13: 비고
"""
import re
from datetime import date, datetime
from pathlib import Path
from dataclasses import dataclass, field
from typing import Optional
import openpyxl


@dataclass
class ComparableSaleRecord:
    source_file: str
    subject_property_serial: str       # 파일명에서 추출
    case_index: int                    # 0=본건, 1~N=거래사례
    address_full: str = ""
    address_sido: str = ""
    address_sigungu: str = ""
    address_dong: str = ""
    property_type: str = ""
    use_zone: str = ""
    land_area: Optional[float] = None
    building_area: Optional[float] = None
    structure: str = ""
    approval_date: Optional[date] = None
    trade_amount: Optional[int] = None
    trade_date: Optional[date] = None
    unit_price_py: Optional[float] = None
    land_price_sqm: Optional[float] = None
    note: str = ""


def _parse_address(addr: str) -> tuple[str, str, str]:
    """'서울특별시 강남구 역삼동 683-7...' → (sido, sigungu, dong)"""
    if not addr:
        return "", "", ""
    parts = str(addr).split()
    sido = parts[0] if len(parts) > 0 else ""
    sigungu = parts[1] if len(parts) > 1 else ""
    dong = parts[2] if len(parts) > 2 else ""
    return sido, sigungu, dong


def _clean_float(v) -> Optional[float]:
    if v is None:
        return None
    try:
        return float(str(v).replace(",", "").replace("-", "").strip() or "0") or None
    except (ValueError, TypeError):
        return None


def _clean_int(v) -> Optional[int]:
    f = _clean_float(v)
    return int(f) if f else None


def _clean_date(v) -> Optional[date]:
    if v is None:
        return None
    if isinstance(v, (date, datetime)):
        return v.date() if isinstance(v, datetime) else v
    try:
        s = str(v).strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(s[:10], fmt).date()
            except ValueError:
                pass
    except Exception:
        pass
    return None


def _extract_serial(file_path: str) -> str:
    """파일명에서 물건번호 추출: 'IBK-R-001 (주)...' → 'IBK-R-001'"""
    stem = Path(file_path).stem
    m = re.match(r"([A-Za-z]+-[A-Za-z0-9]+(?:-[0-9]+(?:_[0-9]+)?)?)", stem)
    return m.group(1) if m else stem[:20]


def parse_comparable_sheet(file_path: str) -> list[ComparableSaleRecord]:
    """정밀 시트를 파싱하여 본건 + 거래사례 레코드 목록 반환"""
    results = []
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    except Exception:
        return results

    if "정밀" not in wb.sheetnames:
        wb.close()
        return results

    ws = wb["정밀"]
    subject_serial = _extract_serial(file_path)

    # 헤더 행(R1) 에서 컬럼 위치 파악
    # R1: [구분, None, None, 본건, None, 거래사례1, None, 거래사례2, ...]
    header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
    case_cols: list[tuple[int, str]] = []  # (0-based col_idx, label)
    for c_idx, val in enumerate(header_row):
        if val is None:
            continue
        s = str(val).strip()
        if s == "본건":
            case_cols.append((c_idx, "본건"))
        elif re.match(r"거래사례\d+", s):
            case_cols.append((c_idx, s))

    if not case_cols:
        wb.close()
        return results

    # 데이터 행 읽기 (R2~R13)
    row_labels = {}  # row_idx(1-based) → 구분 값
    data_rows = {}   # row_idx → [값들]
    for r_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=15, values_only=True), start=2):
        label = str(row[0]).strip() if row[0] is not None else ""
        row_labels[r_idx] = label
        data_rows[r_idx] = list(row)

    # 행 레이블 → row_idx 매핑
    label_to_row = {}
    for r_idx, label in row_labels.items():
        if label:
            label_to_row[label] = r_idx

    def get_cell(label_key, col_idx):
        r = label_to_row.get(label_key)
        if r is None:
            return None
        row = data_rows.get(r, [])
        return row[col_idx] if col_idx < len(row) else None

    for case_num, (col_idx, label) in enumerate(case_cols):
        is_subject = (label == "본건")
        addr = str(get_cell("소재지", col_idx) or "").strip()
        sido, sigungu, dong = _parse_address(addr)

        rec = ComparableSaleRecord(
            source_file=file_path,
            subject_property_serial=subject_serial,
            case_index=0 if is_subject else case_num,
            address_full=addr,
            address_sido=sido,
            address_sigungu=sigungu,
            address_dong=dong,
            property_type=str(get_cell("주용도", col_idx) or "").strip(),
            use_zone=str(get_cell("용도지역", col_idx) or "").strip(),
            land_area=_clean_float(get_cell("토지(㎡)", col_idx)),
            building_area=_clean_float(get_cell("건물(㎡)", col_idx)),
            structure=str(get_cell("주구조", col_idx) or "").strip(),
            approval_date=_clean_date(get_cell("사용승인일", col_idx)),
            trade_amount=_clean_int(get_cell("거래금액", col_idx)),
            trade_date=_clean_date(get_cell("거래시점/기준시점", col_idx)),
            unit_price_py=_clean_float(get_cell("일괄단가(원/py)", col_idx)),
            land_price_sqm=_clean_float(get_cell("개별지가(원/㎡)", col_idx)),
            note=str(get_cell("비고", col_idx) or "").strip(),
        )

        # 주소 또는 거래금액이 없으면 스킵 (빈 사례 컬럼)
        if not addr and not rec.trade_amount:
            continue

        results.append(rec)

    wb.close()
    return results
