import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from pathlib import Path

files = list(Path(r"D:\NPL전례\★ 과거 NPL 데이터").rglob("*SH*Datadisk*Clean*.xlsx"))
if not files:
    files = list(Path(r"D:\NPL전례\★ 과거 NPL 데이터").rglob("*SH*Datadisk*.xlsx"))

f = files[0]
print(f"파일: {f.name}")
wb = openpyxl.load_workbook(str(f), data_only=True, read_only=True)
ws = wb["Sheet C-1(물건정보)"]

# 헤더 행 탐색 (row 7~13)
print("\n=== 헤더 행 (비어있지 않은 셀) ===")
for r_idx, row in enumerate(ws.iter_rows(min_row=7, max_row=13, values_only=True), start=7):
    for c_idx, cell in enumerate(row):
        if cell is not None and str(cell).strip():
            print(f"  R{r_idx}C{c_idx+1}: {cell}")

# 첫 데이터 행
print("\n=== 첫 데이터 행 ===")
for r_idx, row in enumerate(ws.iter_rows(min_row=14, max_row=18, values_only=True), start=14):
    if any(row):
        for c_idx, cell in enumerate(row[:30]):
            if cell is not None:
                print(f"  C{c_idx+1}: {cell}")
        break

wb.close()
