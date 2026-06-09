import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.path.insert(0, '.')
import openpyxl
from app.parsers.ibk_alloc_parser import IBKAllocParser

ibk_file = r"D:\NPL 폴더\IBK 2026-2 Program\2026-2P 우리금융(IBK)\★ IBK 2026-2P A Pre Allocation.xlsx"
wb = openpyxl.load_workbook(ibk_file, data_only=True, read_only=True)
ws = wb["4. 담보(물건)"]

# col_map 디버깅
p = IBKAllocParser()
col_map = p._build_col_map(ws)
print("=== IBK Alloc col_map ===")
for k in ['serial_no','debtor_serial','property_serial','address_sido','land_area',
          'appraisal_type','appraisal_date','appraiser','land_value','building_value',
          'machine_value','outside_value','total_value','is_filed','final_result','hammer_price']:
    v = col_map.get(k)
    if v is not None:
        row9 = list(ws.iter_rows(min_row=9, max_row=9, values_only=True))[0]
        header = row9[v] if v < len(row9) else "?"
        print(f"  {k} = col{v+1}: {header}")

# 감정가 있는 첫 번째 물건
print("\n=== 감정가 있는 물건 확인 ===")
data_start = p._find_data_start(ws)
count = 0
for row in ws.iter_rows(min_row=data_start, values_only=True):
    if not any(row):
        continue
    tv_idx = col_map.get("total_value")
    lv_idx = col_map.get("land_value")
    tv = row[tv_idx] if tv_idx and tv_idx < len(row) else None
    lv = row[lv_idx] if lv_idx and lv_idx < len(row) else None
    if tv and tv != 0:
        sido_idx = col_map.get("address_sido")
        sido = row[sido_idx] if sido_idx else "?"
        print(f"  sido={sido}, land_val={lv}, total_val={tv}")
        count += 1
    if count >= 3:
        break

wb.close()
