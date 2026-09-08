"""scripts/export_comparable_sales_excel.py 검증.

실제 'Loan4U 업로드 엑셀' 템플릿은 이 저장소에 없어 그 정확한 포맷을
재현하지는 못한다 — 대신 loan4u_phase12_pipeline.py 의 실제 가격
부합성 판정 로직(classify_price_conformity_phase12)을 그대로 재사용해
수집 데이터를 등급 매긴 엑셀로 만드는 부분만 검증한다.
"""

import sys
from datetime import date
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))
sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))

from export_comparable_sales_excel import export


def _row(sido, sigungu, addr, ptype, area, amount, tdate=date(2025, 12, 1)):
    return {
        "address_sido": sido, "address_sigungu": sigungu, "address_full": addr,
        "property_type": ptype, "building_area": area, "trade_amount": amount,
        "trade_date": tdate,
    }


class TestExportComparableSalesExcel:
    def test_creates_one_sheet_per_property_type(self, tmp_path):
        rows = [
            _row("서울", "강남구", "A", "아파트", 84.0, 1_200_000_000),
            _row("서울", "송파구", "D", "상가", 50.0, 800_000_000),
        ]
        out = tmp_path / "out.xlsx"
        count = export(rows, out)

        assert count == 2
        from openpyxl import load_workbook
        wb = load_workbook(out)
        assert set(wb.sheetnames) == {"아파트", "상가"}

    def test_price_outlier_is_flagged_against_group_average(self, tmp_path):
        """세 건 중 가격이 튀는 한 건이 실제로 '편차주의'로 잡히는지 —
        loan4u_phase12_pipeline 의 판정 로직을 제대로 재사용했는지의 핵심."""
        rows = [
            _row("서울", "강남구", "A", "아파트", 84.0, 1_200_000_000),
            _row("서울", "강남구", "B", "아파트", 84.0, 1_250_000_000),
            _row("서울", "강남구", "이상치", "아파트", 84.0, 2_400_000_000),
        ]
        out = tmp_path / "out.xlsx"
        export(rows, out)

        from openpyxl import load_workbook
        ws = load_workbook(out)["아파트"]
        grades = {r[2]: r[8] for r in ws.iter_rows(min_row=2, values_only=True)}
        assert grades["이상치"] == "편차주의"

    def test_single_transaction_group_defaults_to_within_range(self, tmp_path):
        """비교 대상이 자신뿐이면(그룹에 1건) 편차 0%로 '적정' 처리된다."""
        rows = [_row("서울", "송파구", "단독", "상가", 50.0, 800_000_000)]
        out = tmp_path / "out.xlsx"
        export(rows, out)

        from openpyxl import load_workbook
        ws = load_workbook(out)["상가"]
        row = list(ws.iter_rows(min_row=2, values_only=True))[0]
        assert row[8] == "적정"

    def test_land_rows_use_land_area_as_unit_price_basis(self, tmp_path):
        """토지는 건물면적이 없다 — 대지면적으로 ㎡당 단가를 내야 '추가확인'으로 빠지지 않는다."""
        rows = [{**_row("경기", "성남시", "밭", "토지", None, 4_500_000_000), "land_area": 1250.5}]
        out = tmp_path / "out.xlsx"
        export(rows, out)

        from openpyxl import load_workbook
        row = list(load_workbook(out)["토지"].iter_rows(min_row=2, values_only=True))[0]
        assert row[6] == pytest.approx(1250.5)
        assert row[7] == round(4_500_000_000 / 1250.5)
        assert row[8] == "적정"

    def test_missing_area_or_amount_does_not_crash(self, tmp_path):
        rows = [
            _row("서울", "강남구", "면적없음", "아파트", None, 1_200_000_000),
            _row("서울", "강남구", "금액없음", "아파트", 84.0, None),
        ]
        out = tmp_path / "out.xlsx"
        count = export(rows, out)
        assert count == 2  # 둘 다 행은 만들어지되 '추가확인'으로 처리

    def test_decimal_trade_amount_from_real_db_does_not_crash(self, tmp_path):
        """실제 DB(app/db/models.ComparableSale.trade_amount)는 Numeric
        컬럼이라 Decimal 로 온다 — float 와 나눌 때 TypeError 가 났던
        실제 버그를 고정한다."""
        from decimal import Decimal
        rows = [_row("서울", "강남구", "A", "아파트", 84.0, Decimal("1200000000"))]
        out = tmp_path / "out.xlsx"
        count = export(rows, out)
        assert count == 1
