"""scripts/collect_all_transactions.py 검증.

네 유형의 API 호출을 가짜 fetch 로 바꿔 오케스트레이션만 본다 — 유형별
집계, 한 유형 실패의 격리, 같은 테이블 적재, 재실행 중복 방지, 엑셀
내보내기. 실제 API 응답 파싱은 각 수집기 테스트가 따로 본다.
"""

import importlib
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))
sys.path.insert(0, str(Path(__file__).parent.parent / "scripts"))

from app.integrations import db_ingest
from app.integrations.korea_api import TransactionRecord
from app.integrations.korea_api_commercial import CommercialTransactionRecord
from app.integrations.korea_api_industrial import IndustrialTransactionRecord
from app.integrations.korea_api_land import LandTransactionRecord
from collect_all_transactions import Collector, collect

INGESTERS = {
    "residential": db_ingest.bulk_ingest_transactions,
    "industrial": db_ingest.bulk_ingest_industrial_transactions,
    "commercial": db_ingest.bulk_ingest_commercial_transactions,
    "land": db_ingest.bulk_ingest_land_transactions,
}


@pytest.fixture
def db(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'all.sqlite'}")
    import app.db.database as database_module
    importlib.reload(database_module)
    database_module.init_db()
    session = database_module.SessionLocal()
    yield session
    session.close()


def _records() -> dict:
    common = dict(sgg_code="11680", contract_year=2025, contract_month=12)
    return {
        "residential": [TransactionRecord(
            property_type="아파트", complex_name="래미안", address_dong="대치동", address_jibun="1",
            contract_day=3, price_manwon=250_000, exclusive_area=84.9, floor=12, **common)],
        "industrial": [IndustrialTransactionRecord(
            building_name="공장A", address_dong="수서동", address_jibun="2",
            contract_day=4, price_manwon=500_000, land_area=2000.0, building_area=1500.0, **common)],
        "commercial": [CommercialTransactionRecord(
            building_name="상가B", address_dong="역삼동", address_jibun="3",
            contract_day=5, price_manwon=820_000, land_area=300.0, building_area=950.0, **common)],
        "land": [LandTransactionRecord(
            address_dong="율현동", address_jibun="4", land_category="전", use_zone="자연녹지지역",
            contract_day=6, price_manwon=450_000, land_area=1250.5, **common)],
    }


def _fail(*_):
    raise RuntimeError("API 최대 재시도 초과")


def _fake_collectors(records: dict, failing: set = frozenset()) -> dict:
    def fetcher(recs):
        return lambda sgg, year, month: recs

    return {
        name: Collector(name, _fail if name in failing else fetcher(records[name]), ingest)
        for name, ingest in INGESTERS.items()
    }


class TestCollectAll:
    def test_all_four_types_land_in_one_table_and_export(self, db, tmp_path):
        results = collect(db, _fake_collectors(_records()), ["11680"], 2025, 12, pause_sec=0)

        assert all(r["inserted"] == 1 and r["failed_calls"] == 0 for r in results.values())

        from app.db.models import ComparableSale
        types = {row.property_type for row in db.query(ComparableSale).all()}
        assert types == {"아파트", "공장창고", "상가", "토지"}

        from export_comparable_sales_excel import export, load_rows
        out = tmp_path / "loan4u.xlsx"
        assert export(load_rows(db), out) == 4
        from openpyxl import load_workbook
        assert set(load_workbook(out).sheetnames) == types

    def test_one_failing_type_does_not_block_others(self, db):
        collectors = _fake_collectors(_records(), failing={"land"})
        results = collect(db, collectors, ["11680"], 2025, 12, pause_sec=0)

        assert results["land"] == {"fetched": 0, "inserted": 0, "skipped": 0, "errors": 0, "failed_calls": 1}
        assert results["commercial"]["inserted"] == 1
        assert results["residential"]["inserted"] == 1

    def test_rerun_dedupes_every_type(self, db):
        collectors = _fake_collectors(_records())
        collect(db, collectors, ["11680"], 2025, 12, pause_sec=0)
        results = collect(db, collectors, ["11680"], 2025, 12, pause_sec=0)

        assert all(r["inserted"] == 0 and r["skipped"] == 1 for r in results.values())

    def test_multiple_sgg_codes_are_all_visited(self, db):
        calls = []

        def fetch(sgg, year, month):
            calls.append(sgg)
            return []

        collectors = {"land": Collector("토지", fetch, db_ingest.bulk_ingest_land_transactions)}
        collect(db, collectors, ["11680", "11650"], 2025, 12, pause_sec=0)
        assert calls == ["11680", "11650"]
