#!/usr/bin/env python3
"""
Loan4U Phase 12 PDF Report Generator
Converts 21-sheet Excel workbook to professional multi-page PDF summary.
"""

import logging
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak,
    KeepTogether
)

log = logging.getLogger(__name__)


@dataclass
class CountryMetric:
    """Country validation metrics"""
    country: str
    total_properties: int
    passed_validation: int
    pass_rate: float
    mape: Optional[float] = None


@dataclass
class SheetData:
    """Data from country sheet"""
    country: str
    sheet_type: str  # 'Residential' or 'Prediction'
    properties: int
    valid_count: int
    headers: List[str]
    rows: List[List]


@dataclass
class ReportMetrics:
    """Global report metrics"""
    review_date: str
    total_sheets: int
    total_records: int
    countries: List[str]
    country_metrics: List[CountryMetric]
    sheet_data: Dict[str, SheetData]


class Phase12PDFGenerator:
    """Generate Phase 12 validation report PDF from 21-sheet Excel."""

    GRADE_COLORS = {
        '적정': colors.HexColor('#C6EFCE'),
        '확인필요': colors.HexColor('#FFEB9C'),
        '편차주의': colors.HexColor('#FFC7CE'),
        '추가확인': colors.HexColor('#FF0000'),
    }

    GRADE_LABELS = {
        '적정': 'Acceptable',
        '확인필요': 'Needs Review',
        '편차주의': 'High Variance',
        '추가확인': 'Action Required',
    }

    def __init__(self, output_path: str) -> None:
        self.output_path = Path(output_path)
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        logging.basicConfig(level=logging.INFO)

    def generate_from_excel(self, excel_path: str, output_pdf: str) -> bool:
        """Generate comprehensive PDF report from Excel workbook."""
        try:
            wb = load_workbook(excel_path, data_only=True)
            if 'Report' not in wb.sheetnames:
                log.error("Report sheet not found")
                return False

            metrics = self._extract_metrics(wb)
            self._create_pdf_report(metrics, output_pdf)
            wb.close()
            log.info(f"✅ PDF generated: {output_pdf}")
            return True
        except Exception as e:
            import traceback
            log.error(f"PDF generation failed: {e}")
            traceback.print_exc()
            return False

    def _extract_metrics(self, wb) -> ReportMetrics:
        """Extract metrics from all sheets."""
        ws = wb['Report']
        review_date = str(ws['A2'].value or datetime.now().strftime("%Y-%m-%d"))

        country_metrics = []
        sheet_data = {}
        total_records = 0

        # Parse Report sheet (rows 6-13 are countries)
        for row in range(6, 14):
            country = ws.cell(row, 1).value
            props = ws.cell(row, 2).value or 0
            passed = ws.cell(row, 3).value or 0

            if country and props:
                total_records += int(props)
                rate = float(passed) / int(props) if props > 0 else 0
                country_metrics.append(
                    CountryMetric(country, int(props), int(passed), rate)
                )

        # Extract detailed data from country sheets
        for country in [m.country for m in country_metrics]:
            for sheet_type in ['Residential', 'Prediction']:
                sheet_name = f"{country}_{sheet_type}"
                if sheet_name in wb.sheetnames:
                    data = self._extract_sheet_data(wb[sheet_name], country, sheet_type)
                    sheet_data[sheet_name] = data

        countries = [m.country for m in country_metrics]

        return ReportMetrics(
            review_date=review_date,
            total_sheets=len(wb.sheetnames),
            total_records=total_records,
            countries=countries,
            country_metrics=country_metrics,
            sheet_data=sheet_data,
        )

    def _extract_sheet_data(self, ws, country: str, sheet_type: str) -> SheetData:
        """Extract data from country-specific sheet."""
        headers = []
        rows = []

        if ws.max_row > 0:
            headers = [ws.cell(1, col).value or '' for col in range(1, ws.max_column + 1)]
            for row in range(2, ws.max_row + 1):
                row_data = [ws.cell(row, col).value or '' for col in range(1, ws.max_column + 1)]
                rows.append(row_data)

        return SheetData(
            country=country,
            sheet_type=sheet_type,
            properties=len(rows),
            valid_count=sum(1 for r in rows if r and str(r[0]).strip()),
            headers=headers,
            rows=rows,
        )

    def _create_pdf_report(self, metrics: ReportMetrics, output_pdf: str) -> None:
        """Create multi-page PDF document."""
        doc = SimpleDocTemplate(
            output_pdf,
            pagesize=letter,
            rightMargin=0.5 * inch,
            leftMargin=0.5 * inch,
            topMargin=0.75 * inch,
            bottomMargin=0.75 * inch,
        )

        styles = self._create_styles()
        story = []

        # Page 1: Title and Summary
        story.extend(self._build_cover_page(metrics, styles))
        story.append(PageBreak())

        # Page 2: Country Summary Table
        story.extend(self._build_summary_page(metrics, styles))
        story.append(PageBreak())

        # Pages 3+: Country Details
        for country in metrics.countries:
            country_content = self._build_country_page(metrics, country, styles)
            if country_content:
                story.extend(country_content)
                story.append(PageBreak())

        # Final page: Footer
        story.extend(self._build_footer_page(metrics, styles))

        doc.build(story)

    def _create_styles(self) -> Dict[str, ParagraphStyle]:
        """Create custom paragraph styles."""
        styles = getSampleStyleSheet()
        custom = {
            'title': ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=20,
                textColor=colors.HexColor('#1F4E78'),
                spaceAfter=12,
                alignment=1,
            ),
            'heading2': ParagraphStyle(
                'CustomHeading2',
                parent=styles['Heading2'],
                fontSize=14,
                textColor=colors.HexColor('#2E5C8A'),
                spaceAfter=10,
            ),
            'heading3': ParagraphStyle(
                'CustomHeading3',
                parent=styles['Heading3'],
                fontSize=12,
                textColor=colors.HexColor('#404040'),
                spaceAfter=8,
            ),
            'normal': styles['Normal'],
        }
        return custom

    def _build_cover_page(self, metrics: ReportMetrics, styles: Dict) -> List:
        """Build cover page with title and summary."""
        elements = []

        # Title
        elements.append(Paragraph("Loan4U Phase 12", styles['title']))
        elements.append(Paragraph("Global Validation Report", styles['title']))
        elements.append(Spacer(1, 0.3 * inch))

        # Review metadata
        elements.append(Paragraph(
            f"<b>Review Date:</b> {metrics.review_date}",
            styles['normal']
        ))
        elements.append(Paragraph(
            f"<b>Generated:</b> {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            styles['normal']
        ))
        elements.append(Spacer(1, 0.3 * inch))

        # Summary statistics
        total_passed = sum(m.passed_validation for m in metrics.country_metrics)
        summary = (
            f"<b>Summary:</b><br/>"
            f"• {metrics.total_sheets} sheets processed<br/>"
            f"• {metrics.total_records} properties validated<br/>"
            f"• {len(metrics.countries)} countries analyzed<br/>"
            f"• {total_passed} properties passed validation"
        )
        elements.append(Paragraph(summary, styles['normal']))
        elements.append(Spacer(1, 0.3 * inch))

        # Overall pass rate
        overall_rate = total_passed / metrics.total_records if metrics.total_records > 0 else 0
        elements.append(Paragraph(
            f"<b>Overall Pass Rate:</b> {overall_rate:.1%}",
            styles['normal']
        ))

        return elements

    def _build_summary_page(self, metrics: ReportMetrics, styles: Dict) -> List:
        """Build country summary table page."""
        elements = []
        elements.append(Paragraph("Country Summary", styles['heading2']))
        elements.append(Spacer(1, 0.2 * inch))

        table_data = [['Country', 'Properties', 'Passed', 'Pass Rate', 'Status']]
        for m in metrics.country_metrics:
            status = self._get_status_label(m.pass_rate)
            table_data.append([
                m.country,
                str(m.total_properties),
                str(m.passed_validation),
                f"{m.pass_rate:.1%}",
                status,
            ])

        table = Table(table_data, colWidths=[1.2 * inch, 1.0 * inch, 1.0 * inch, 1.0 * inch, 1.3 * inch])
        table.setStyle(self._get_table_style())
        elements.append(table)

        return elements

    def _build_country_page(self, metrics: ReportMetrics, country: str, styles: Dict) -> Optional[List]:
        """Build country detail page."""
        elements = []
        elements.append(Paragraph(f"{country} Validation Results", styles['heading2']))
        elements.append(Spacer(1, 0.2 * inch))

        # Country metrics
        country_metric = next((m for m in metrics.country_metrics if m.country == country), None)
        if country_metric:
            metric_text = (
                f"<b>Total Properties:</b> {country_metric.total_properties} | "
                f"<b>Passed:</b> {country_metric.passed_validation} | "
                f"<b>Pass Rate:</b> {country_metric.pass_rate:.1%}"
            )
            elements.append(Paragraph(metric_text, styles['normal']))
            elements.append(Spacer(1, 0.15 * inch))

        # Residential and Prediction data
        for sheet_type in ['Residential', 'Prediction']:
            sheet_name = f"{country}_{sheet_type}"
            if sheet_name in metrics.sheet_data:
                data = metrics.sheet_data[sheet_name]
                elements.append(Paragraph(f"{sheet_type} Data", styles['heading3']))

                if data.headers and data.rows:
                    table_data = [data.headers[:5]]  # Limit to first 5 columns
                    for row in data.rows[:3]:  # Limit to first 3 data rows
                        table_data.append([str(cell)[:20] for cell in row[:5]])

                    table = Table(table_data, colWidths=[1.5 * inch] * 5)
                    table.setStyle(self._get_data_table_style())
                    elements.append(table)
                elements.append(Spacer(1, 0.2 * inch))

        return elements if len(elements) > 1 else None

    def _build_footer_page(self, metrics: ReportMetrics, styles: Dict) -> List:
        """Build footer with summary and notes."""
        elements = []
        elements.append(Paragraph("Report Summary", styles['heading2']))
        elements.append(Spacer(1, 0.15 * inch))

        total_passed = sum(m.passed_validation for m in metrics.country_metrics)
        quality_rate = total_passed / metrics.total_records if metrics.total_records > 0 else 0
        quality_label = '✓ Excellent' if quality_rate > 0.95 else '⚠ Acceptable' if quality_rate > 0.7 else '✗ Needs Attention'

        summary = (
            f"<b>Validation Completion:</b> {total_passed}/{metrics.total_records} properties passed<br/>"
            f"<b>Quality Level:</b> {quality_label}<br/>"
            f"<b>Next Steps:</b> Review properties marked as 'Action Required' and 'High Variance'"
        )
        elements.append(Paragraph(summary, styles['normal']))
        elements.append(Spacer(1, 0.3 * inch))

        footer = (
            f"<i>Report generated {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | "
            f"Loan4U Phase 12 Validation | "
            f"See .claude/EXECUTION_POLICY.md for policies</i>"
        )
        elements.append(Paragraph(footer, styles['normal']))

        return elements

    def _get_status_label(self, pass_rate: float) -> str:
        """Get status label based on pass rate."""
        if pass_rate >= 0.95:
            return "✓ Excellent"
        elif pass_rate >= 0.80:
            return "✓ Good"
        elif pass_rate >= 0.60:
            return "⚠ Fair"
        else:
            return "✗ Poor"

    def _get_table_style(self) -> TableStyle:
        """Return main table styling."""
        return TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 11),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
            ('TOPPADDING', (0, 0), (-1, 0), 10),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#CCCCCC')),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5F5')]),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ])

    def _get_data_table_style(self) -> TableStyle:
        """Return data table styling."""
        return TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2E5C8A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('TOPPADDING', (0, 0), (-1, 0), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#DDDDDD')),
            ('FONTNAME', (0, 1), (-1, -1), 'Courier'),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FAFAFA')]),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ])


def main() -> None:
    """Generate PDF from Phase 12 output Excel."""
    import argparse
    import glob

    parser = argparse.ArgumentParser(description='Phase 12 PDF Report Generator')
    parser.add_argument(
        '--excel',
        default=None,
        help='Input Excel workbook path (auto-detect if omitted)',
    )
    parser.add_argument(
        '--output',
        default='output/Loan4U_Phase12_Final_Report.pdf',
        help='Output PDF path',
    )

    args = parser.parse_args()

    # Auto-detect Excel file if not specified
    excel_path = args.excel
    if not excel_path:
        excel_files = glob.glob('output/Loan4U*.xlsx')
        if not excel_files:
            print("✗ No Excel file found in output/")
            exit(1)
        excel_path = sorted(excel_files)[-1]
        print(f"Auto-detected: {excel_path}")

    generator = Phase12PDFGenerator(args.output)
    success = generator.generate_from_excel(excel_path, args.output)

    if success:
        print(f"✅ PDF generated: {args.output}")
    else:
        print(f"❌ PDF generation failed")
        exit(1)


if __name__ == '__main__':
    main()
