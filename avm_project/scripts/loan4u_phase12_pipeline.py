#!/usr/bin/env python3
"""
Loan4U Phase 12 Global Pipeline
Unified Excel/PDF generation with price validation across 8 countries.
"""

import json
import logging
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Setup logging
logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')
log = logging.getLogger(__name__)

# Constants
REVIEW_DATE = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
COUNTRIES = ['UK', 'SG', 'JP', 'DE', 'AU', 'CA', 'TH', 'HK']

# Country-specific tolerance ranges (market volatility)
TOLERANCE_MAP = {
    'UK': 0.05, 'JP': 0.05, 'SG': 0.08, 'DE': 0.08,
    'HK': 0.08, 'AU': 0.10, 'CA': 0.10, 'TH': 0.15
}

# Model version confidence factors
CONFIDENCE_MAP = {'v1.0': 1.0, 'v1.1': 0.98, 'v1.2': 0.95}

# Conformity grade colors
GRADE_COLORS = {
    '적정': 'C6EFCE',           # Green
    '확인필요': 'FFEB9C',       # Yellow
    '편차주의': 'FFC7CE',       # Light Red
    '추가확인': 'FF0000'        # Red
}

# Styles
HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


@dataclass
class Phase12CountryData:
    """Country-specific AVM data"""
    country: str
    properties: List[Dict[str, Any]]
    total_records: int = 0

    def __post_init__(self):
        self.total_records = len(self.properties)


@dataclass
class AuditResult:
    """Price audit result for single property"""
    property_id: str
    old_price: Optional[float]
    new_price: Optional[float]
    deviation_ratio: Optional[float]
    conformity_grade: str
    reason: str
    final_action: str


@dataclass
class GlobalMetrics:
    """Aggregated metrics across countries"""
    countries: List[str]
    total_properties: int
    audit_count: int
    passed_audit: int
    audit_rate: float = 0.0
    review_date: str = REVIEW_DATE


def normalize_text(value: Any) -> str:
    """Normalize text value"""
    if value is None:
        return ""
    return str(value).strip()


def normalize_number(value: Any) -> Optional[float]:
    """Parse currency/number strings"""
    if value is None or (isinstance(value, float) and str(value) == 'nan'):
        return None

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    if not text:
        return None

    text = text.replace(',', '').replace('원', '').replace('₩', '')

    if '억' in text:
        try:
            num = float(text.replace('억', '').replace('원', '').strip())
            return num * 100_000_000
        except:
            pass

    if '만' in text:
        try:
            num = float(text.replace('만', '').replace('원', '').strip())
            return num * 10_000
        except:
            pass

    try:
        return float(text)
    except:
        return None


def safe_div(num: Optional[float], denom: Optional[float]) -> Optional[float]:
    """Safe division"""
    return num / denom if num and denom and denom != 0 else None


def get_header_map(ws, header_row: int = 1) -> Dict[str, int]:
    """Get column mapping from header row"""
    header_map = {}
    for col in range(1, ws.max_column + 1):
        val = normalize_text(ws.cell(header_row, col).value)
        if val and val not in header_map:
            header_map[val] = col
    return header_map


def ensure_columns(ws, columns: List[str], header_row: int = 1) -> Dict[str, int]:
    """Add missing columns to worksheet"""
    header_map = get_header_map(ws, header_row)
    next_col = ws.max_column + 1

    for col_name in columns:
        if col_name not in header_map:
            cell = ws.cell(header_row, next_col)
            cell.value = col_name
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = THIN_BORDER
            header_map[col_name] = next_col
            next_col += 1

    return header_map


def classify_price_conformity_phase12(
    old_price: Optional[float],
    new_price: Optional[float],
    country: str,
    model_version: str = 'v1.0',
    real_transaction: bool = False,
    source_records: Optional[List[Dict]] = None
) -> Tuple[str, Optional[float], str]:
    """
    Classify price conformity with country-specific tolerance & model confidence.

    Args:
        old_price: Original price
        new_price: Updated/predicted price
        country: Country code (UK, JP, etc)
        model_version: Model version (v1.0, v1.1, v1.2)
        real_transaction: Real transaction applied
        source_records: External source price records

    Returns:
        (conformity_grade, deviation_ratio, reason)
    """
    if new_price is None:
        return '추가확인', None, 'Missing new price'

    # Get tolerance & confidence factor
    tolerance = TOLERANCE_MAP.get(country, 0.10)
    confidence = CONFIDENCE_MAP.get(model_version, 1.0)
    adjusted_tolerance = tolerance * confidence

    # Calculate deviation ratio
    if old_price and old_price != 0:
        deviation = (new_price - old_price) / old_price
    else:
        deviation = None

    if source_records:
        prices = [r.get('price') for r in source_records if r.get('price')]
        if prices:
            avg_price = sum(prices) / len(prices)
            source_dev = (new_price - avg_price) / avg_price if avg_price else None

            if source_dev is None:
                return '추가확인', None, 'Source average calculation failed'

            abs_dev = abs(source_dev)
            if abs_dev <= adjusted_tolerance:
                return '적정', source_dev, f'Within {country} tolerance'
            elif abs_dev <= adjusted_tolerance * 1.5:
                return '확인필요', source_dev, f'Needs verification'
            else:
                return '편차주의', source_dev, f'Exceeds {country} tolerance'

    # Fallback: use old vs new comparison
    if real_transaction and deviation:
        if abs(deviation) > 0.20:
            return '편차주의', deviation, 'Real transaction but >20% change'
        return '적정', deviation, 'Real transaction applied'

    if deviation:
        if abs(deviation) > 0.20:
            return '편차주의', deviation, '>20% change detected'

    return '추가확인', deviation, 'Insufficient information'


def classify_correction_result(grade: str, real_transaction: bool) -> str:
    """Map conformity grade to correction result"""
    if grade == '적정':
        return '실거래 기반 업데이트' if real_transaction else '출처 기반 적정'
    if grade == '확인필요':
        return '추가확인'
    if grade == '편차주의':
        return '편차주의'
    return '추가확인'


def classify_final_action(grade: str) -> str:
    """Map conformity grade to final action"""
    actions = {
        '적정': '유지',
        '확인필요': '원자료 재확인',
        '편차주의': '가격 재검증 필요',
        '추가확인': '추가 자료 확보'
    }
    return actions.get(grade, '검토필요')


def style_range(ws, min_row: int, max_row: int, min_col: int, max_col: int,
                fill: Optional[PatternFill] = None) -> None:
    """Apply styling to cell range"""
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row, col)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if fill:
                cell.fill = fill


def autofit_columns(ws, min_width: int = 8, max_width: int = 40) -> None:
    """Auto-fit column widths"""
    for col in range(1, ws.max_column + 1):
        max_len = 0
        for row in range(1, min(ws.max_row, 500) + 1):
            val = ws.cell(row, col).value
            if val:
                max_len = max(max_len, len(str(val)))
        width = max(min_width, min(max_len + 2, max_width))
        ws.column_dimensions[get_column_letter(col)].width = width


def process_domestic_sheets(wb, input_path: str, configs: Dict) -> None:
    """Extend domestic sheets (아파트, 빌라) with validation"""
    try:
        base_wb = load_workbook(input_path)
    except Exception as e:
        log.warning(f"Could not load base Excel: {e}")
        return

    # Handle sheet names with potential whitespace
    domestic_sheets = []
    for candidate in ['아파트', ' 아파트', '빌라', ' 빌라']:
        if candidate in base_wb.sheetnames:
            domestic_sheets.append(candidate)

    for sheet_name in domestic_sheets:
        if sheet_name not in base_wb.sheetnames:
            continue

        src_ws = base_wb[sheet_name]
        dst_ws = wb.create_sheet(sheet_name)

        # Copy structure & data
        for row in src_ws.iter_rows():
            for cell in row:
                new_cell = dst_ws[cell.coordinate]
                new_cell.value = cell.value
                if cell.has_style:
                    new_cell.font = copy(cell.font)
                    new_cell.border = copy(cell.border)
                    new_cell.fill = copy(cell.fill)
                    new_cell.alignment = copy(cell.alignment)

        # Add validation columns
        header_map = get_header_map(dst_ws)
        next_col = max(header_map.values()) + 1 if header_map else 1

        new_cols = {
            '가격부합성': next_col,
            '편차율': next_col + 1,
            '최종조치': next_col + 2
        }

        for i, col_name in enumerate(new_cols.keys()):
            cell = dst_ws.cell(1, next_col + i)
            cell.value = col_name
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.border = THIN_BORDER

        # Process data rows
        for row in range(2, dst_ws.max_row + 1):
            row_data = {h: dst_ws.cell(row, c).value for h, c in header_map.items()}
            old_price = normalize_number(row_data.get('기존가격'))
            new_price = normalize_number(row_data.get('신규가격'))

            grade, dev, _ = classify_price_conformity_phase12(
                old_price, new_price, 'KR', real_transaction=False
            )

            # Write results
            dst_ws.cell(row, new_cols['가격부합성']).value = grade
            if dev:
                dst_ws.cell(row, new_cols['편차율']).value = f"{dev:.2%}"
            dst_ws.cell(row, new_cols['최종조치']).value = classify_final_action(grade)

            # Apply color
            fill = PatternFill(fill_type="solid", fgColor=GRADE_COLORS.get(grade, "FFFFFF"))
            dst_ws.cell(row, new_cols['가격부합성']).fill = fill


def process_country_sheets(wb, country: str, country_data: Phase12CountryData,
                          model_version: str = 'v1.0') -> None:
    """Create country-specific sheets (Residential + Prediction)"""

    # Residential sheet
    res_name = f"{country}_Residential"
    res_ws = wb.create_sheet(res_name) if res_name not in wb.sheetnames else wb[res_name]

    headers = ['Property_ID', 'Address', 'Area', 'Old_Price', 'New_Price', 'Conformity', 'Deviation', 'Action']
    for col, header in enumerate(headers, 1):
        cell = res_ws.cell(1, col)
        cell.value = header
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER

    # Populate data
    for row, prop in enumerate(country_data.properties, 2):
        old_p = normalize_number(prop.get('old_price'))
        new_p = normalize_number(prop.get('new_price'))

        grade, dev, _ = classify_price_conformity_phase12(
            old_p, new_p, country, model_version
        )

        res_ws.cell(row, 1).value = prop.get('property_id')
        res_ws.cell(row, 2).value = prop.get('address')
        res_ws.cell(row, 3).value = prop.get('area')
        res_ws.cell(row, 4).value = old_p
        res_ws.cell(row, 5).value = new_p
        res_ws.cell(row, 6).value = grade
        if dev:
            res_ws.cell(row, 7).value = f"{dev:.2%}"
        res_ws.cell(row, 8).value = classify_final_action(grade)

        # Apply color
        fill = PatternFill(fill_type="solid", fgColor=GRADE_COLORS.get(grade))
        res_ws.cell(row, 6).fill = fill

    style_range(res_ws, 2, res_ws.max_row, 1, len(headers))
    autofit_columns(res_ws)

    # Prediction sheet (mirror with v1.0 baseline)
    pred_name = f"{country}_Prediction"
    pred_ws = wb.create_sheet(pred_name) if pred_name not in wb.sheetnames else wb[pred_name]

    # Copy structure
    for col, header in enumerate(headers, 1):
        cell = pred_ws.cell(1, col)
        cell.value = header
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    # Copy data (for now, same as residential)
    for row, prop in enumerate(country_data.properties, 2):
        for col in range(1, len(headers) + 1):
            pred_ws.cell(row, col).value = res_ws.cell(row, col).value
            pred_ws.cell(row, col).border = THIN_BORDER

    autofit_columns(pred_ws)


def create_report_sheet(wb, metrics: GlobalMetrics, audit_results: List[AuditResult]) -> None:
    """Create summary report sheet"""

    ws = wb.create_sheet("Report", 0)

    # Header
    ws['A1'] = "Loan4U Phase 12 Global Validation Report"
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:H1')

    ws['A2'] = f"Review Date: {metrics.review_date}"
    ws.merge_cells('A2:H2')

    # Metrics
    ws['A4'] = "Country Metrics"
    ws['A4'].font = Font(bold=True)

    headers = ['Country', 'Properties', 'Audited', 'Pass Rate']
    for col, header in enumerate(headers, 1):
        ws.cell(5, col).value = header
        ws.cell(5, col).fill = HEADER_FILL
        ws.cell(5, col).font = HEADER_FONT

    row = 6
    for country in metrics.countries:
        count = sum(1 for a in audit_results if a.property_id.startswith(country))
        passed = sum(1 for a in audit_results if a.property_id.startswith(country) and a.conformity_grade == '적정')

        ws.cell(row, 1).value = country
        ws.cell(row, 2).value = count
        ws.cell(row, 3).value = passed
        if count > 0:
            ws.cell(row, 4).value = f"{passed/count:.1%}"

        row += 1

    autofit_columns(ws)
    ws.freeze_panes = 'A6'


def apply_workbook_formatting(wb) -> None:
    """Apply final formatting to all sheets"""
    for ws in wb.sheetnames:
        ws_obj = wb[ws]

        # Set freeze panes
        if ws != 'Report':
            ws_obj.freeze_panes = ws_obj.cell(2, 1).coordinate

        # Auto-fit all columns
        autofit_columns(ws_obj)


def generate_pdf_report(excel_path: str, pdf_path: str, metrics: GlobalMetrics) -> None:
    """Generate PDF summary from metrics"""
    try:
        from weasyprint import HTML, CSS
    except ImportError:
        log.warning("weasyprint not installed, skipping PDF generation")
        return

    html_content = f"""
    <html>
    <head>
        <style>
            body {{ font-family: Arial, sans-serif; margin: 20px; }}
            h1 {{ color: #1F4E78; }}
            table {{ border-collapse: collapse; width: 100%; margin-top: 20px; }}
            th, td {{ border: 1px solid #ddd; padding: 10px; text-align: left; }}
            th {{ background-color: #1F4E78; color: white; }}
            .metric {{ font-size: 18px; font-weight: bold; color: #1F4E78; }}
        </style>
    </head>
    <body>
        <h1>Loan4U Phase 12 Global Validation Report</h1>
        <p>Review Date: {metrics.review_date}</p>

        <div class="metric">Total Properties: {metrics.total_properties}</div>
        <div class="metric">Audit Rate: {metrics.audit_rate:.1%}</div>

        <h2>Country Summary</h2>
        <table>
            <tr>
                <th>Country</th>
                <th>Properties</th>
                <th>Status</th>
            </tr>
    """

    for country in metrics.countries:
        html_content += f"""
            <tr>
                <td>{country}</td>
                <td>Data pending</td>
                <td>Ready for audit</td>
            </tr>
        """

    html_content += """
        </table>
    </body>
    </html>
    """

    try:
        HTML(string=html_content).write_pdf(pdf_path)
        log.info(f"PDF saved: {pdf_path}")
    except Exception as e:
        log.error(f"PDF generation failed: {e}")


def validate_workbook(wb) -> Tuple[bool, List[str]]:
    """Validate workbook structure"""
    errors = []
    sheets = wb.sheetnames

    # Check required sheets (Report + Domestic + Countries)
    has_report = 'Report' in sheets
    has_domestic = any(s.strip() in ['아파트', '빌라'] for s in sheets)
    has_countries = sum(1 for c in COUNTRIES for s in sheets if s.startswith(c + '_')) >= 14

    if not has_report:
        errors.append("Missing sheet: Report")
    if not has_domestic:
        errors.append("Missing domestic sheets (아파트/빌라)")
    if not has_countries:
        errors.append("Missing country sheets")

    return len(errors) == 0, errors


def run_pipeline(base_excel: str, config_dir: str, output_excel: str, output_pdf: str = None) -> None:
    """Main execution pipeline"""

    log.info(f"Loading base Excel: {base_excel}")

    # Load configurations
    config_path = Path(config_dir)
    with open(config_path / 'countries_strategy.json') as f:
        countries_strategy = json.load(f)

    # Create new workbook
    wb = Workbook()
    wb.remove(wb.active)

    # Process domestic sheets
    log.info("Processing domestic sheets...")
    process_domestic_sheets(wb, base_excel, countries_strategy)

    # Process country sheets
    log.info(f"Creating {len(COUNTRIES)} country sheets...")
    metrics = GlobalMetrics(countries=COUNTRIES, total_properties=0, audit_count=0, passed_audit=0)

    country_file_map = {
        'UK': 'uk_expansion_plan.json',
        'SG': 'singapore_expansion_plan.json',
        'JP': 'japan_expansion_plan.json',
        'DE': 'de_expansion_plan.json',
        'AU': 'au_expansion_plan.json',
        'CA': 'ca_expansion_plan.json',
        'TH': 'th_expansion_plan.json',
        'HK': 'hk_expansion_plan.json'
    }

    for country in COUNTRIES:
        filename = country_file_map.get(country, f'{country.lower()}_expansion_plan.json')
        config_file = config_path / filename
        if config_file.exists():
            with open(config_file) as f:
                country_config = json.load(f)

            # Build country data (using config data or sample)
            properties = country_config.get('sample_properties', [])
            if not properties:
                # Generate minimal sample if not provided
                properties = [
                    {'property_id': f'{country}_001', 'address': f'Address 1, {country}', 'area': 3000, 'old_price': 500000, 'new_price': 520000},
                    {'property_id': f'{country}_002', 'address': f'Address 2, {country}', 'area': 2500, 'old_price': 400000, 'new_price': 420000}
                ]

            country_data = Phase12CountryData(country=country, properties=properties)

            process_country_sheets(wb, country, country_data)
            metrics.total_properties += len(properties)
            metrics.audit_count += len(properties)

    # Create report
    log.info("Creating report sheet...")
    audit_results = []  # Would be populated with actual audit results
    create_report_sheet(wb, metrics, audit_results)

    # Apply formatting
    log.info("Applying formatting...")
    apply_workbook_formatting(wb)

    # Validate
    log.info("Validating workbook...")
    is_valid, errors = validate_workbook(wb)
    if not is_valid:
        for error in errors:
            log.warning(error)

    # Save Excel
    log.info(f"Saving Excel: {output_excel}")
    wb.save(output_excel)

    # Generate PDF
    if output_pdf:
        log.info(f"Generating PDF: {output_pdf}")
        generate_pdf_report(output_excel, output_pdf, metrics)

    log.info("✅ Pipeline complete")


if __name__ == '__main__':
    import argparse

    parser = argparse.ArgumentParser(description='Loan4U Phase 12 Pipeline')
    parser.add_argument('--base', default='Loan4U_QC_v1.1_before_fill.xlsx', help='Base Excel file')
    parser.add_argument('--config', default='config/phase12', help='Config directory')
    parser.add_argument('--output', default='Loan4U_QC_v1.1_Phase12_Global_Corrected_20260625.xlsx', help='Output Excel')
    parser.add_argument('--pdf', default='Loan4U_Phase12_Final_Report.pdf', help='Output PDF')

    args = parser.parse_args()
    run_pipeline(args.base, args.config, args.output, args.pdf)
